"""Extract the workbook values the Control Room app displays into data/control_room_workbook.json.
Values are read as calculated by Excel (cached), so the app shows exactly what the workbook shows.
Run: python tools/extract_app_data.py"""
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
wb = load_workbook(ROOT / "Hotel_Food_Waste_Model.xlsx", data_only=True)

def table(sheet, first_row, cols, stop_blank=True, max_row=None):
    ws = wb[sheet]; out = []
    for r in range(first_row, (max_row or ws.max_row) + 1):
        vals = [ws.cell(r, c).value for c in cols]
        if stop_blank and vals[0] in (None, ""): break
        out.append(vals)
    return out

W = wb["Waste Cost"]; E = wb["Economics"]
data = {
    "_about": "Extracted from Hotel_Food_Waste_Model.xlsx (modeled 12-week pilot scenario, not realized hotel performance).",
    "weekly_waste": [{"week": a, "phase": b, "kg": c} for a, b, c in table("Weekly Waste", 2, [1, 2, 3])],
    "waste_by_source": [{"source": a, "baseline": b, "with_agents": c, "change": d} for a, b, c, d in table("Waste by Source", 2, [1, 2, 3, 4])],
    "forecast_mape": [{"week": a, "phase": b, "mape": c} for a, b, c in table("Forecast MAPE", 2, [1, 2, 3])],
    "guest": [{"week": a, "phase": b, "waste_per_cover": c, "guest_score": d} for a, b, c, d in table("Guest Experience", 2, [1, 2, 3, 4], max_row=13)],
    "coordination_hours": [{"week": a, "hours": b} for a, b in table("Operations", 2, [1, 2], max_row=10)],
    "waste_cost_steps": [{"step": a, "amount": b} for a, b in table("Waste Cost", 2, [1, 2], max_row=7)],
    "waste_cost": {"baseline": W["G3"].value, "annual_purchasing": W["F3"].value if isinstance(W["F3"].value, (int, float)) else W["G2"].value,
                   "avoided": W["B9"].value, "reduction": W["B10"].value},
    "agentops_weekly": [dict(zip(["week", "recommendations", "accepted", "overrides", "required_escalations", "correct_escalations",
                                  "all_escalations", "tool_calls", "successful_tool_calls", "acceptance", "override_rate", "recall"], row))
                        for row in table("AgentOps", 2, list(range(1, 13)), max_row=9)],
    "overrides": [{"reason": a, "count": b, "share": c, "improvement": d} for a, b, c, d in table("Overrides", 2, [1, 2, 3, 4], max_row=8)],
    "control_room": [dict(zip(["category", "kpi", "value", "test", "target", "status", "source"], r)) for r in table("Control Room", 2, list(range(1, 8)), max_row=16)],
    "economics": {"food_cost_avoided": E["B3"].value, "coordination_value": E["B4"].value, "gross_benefit": E["B5"].value,
                  "cost_lines": {E[f"A{r}"].value: E[f"B{r}"].value for r in (7, 8, 9, 10)}, "annual_ai_cost": E["B11"].value,
                  "one_time_lines": {E[f"A{r}"].value: E[f"B{r}"].value for r in (13, 14)}, "one_time_total": E["B15"].value,
                  "manager_rate": E["B19"].value, "hours_saved_week": E["B20"].value, "review_hours_week": E["B21"].value,
                  "net_annual": E["B24"].value, "year1_net": E["B25"].value, "year1_roi": E["B26"].value, "payback_months": E["B27"].value,
                  "value_per_ai_cost": E["B28"].value},
    "sensitivity_cases": {c: {"waste_reduction": wb["Sensitivity"][f"{col}2"].value, "adoption": wb["Sensitivity"][f"{col}3"].value,
                              "cost_multiplier": wb["Sensitivity"][f"{col}4"].value, "volume_index": wb["Sensitivity"][f"{col}5"].value,
                              "net_annual": wb["Sensitivity"][f"{col}12"].value, "payback_months": wb["Sensitivity"][f"{col}13"].value}
                          for c, col in (("Conservative", "B"), ("Base", "C"), ("Upside", "D"))},
}
(ROOT / "data" / "control_room_workbook.json").write_text(json.dumps(data, indent=1, default=str))
print("wrote data/control_room_workbook.json")
